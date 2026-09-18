# -*- coding: utf-8 -*-
"""
VSP 사업계획 발표자료(.pptx) — 한국어본 + 베트남어본 두 파일 생성.
아티팩트 HTML 데크(15슬라이드, AI 자동등록 포함)와 같은 내용.

  python scripts/make-decks.py
출력: docs/VSP_사업계획서_KO.pptx , docs/VSP_사업계획서_VI.pptx
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# 팔레트 (옻칠 레드 계열 — 웹 데크와 맞춤)
INK       = RGBColor(0x1B, 0x1A, 0x17)
INK_SOFT  = RGBColor(0x4C, 0x48, 0x42)
INK_FAINT = RGBColor(0x84, 0x7C, 0x70)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
PAPER     = RGBColor(0xF3, 0xEF, 0xE7)
PAPER2    = RGBColor(0xEE, 0xE8, 0xDD)
LINE      = RGBColor(0xD9, 0xD1, 0xC2)
ACCENT    = RGBColor(0xB5, 0x34, 0x1F)   # lacquer
ACC_SOFT  = RGBColor(0xF4, 0xDD, 0xD6)
JADE      = RGBColor(0x1E, 0x6E, 0x5C)
JADE_SOFT = RGBColor(0xDC, 0xEA, 0xE4)
GOLD      = RGBColor(0x9A, 0x72, 0x12)
GOLD_SOFT = RGBColor(0xF0, 0xE7, 0xCE)
COVER_BG  = RGBColor(0x1B, 0x1A, 0x17)

FONT, FONT_M = "Be Vietnam Pro", "Consolas"
W, H = Inches(13.333), Inches(7.5)
ML, MR = Inches(0.85), Inches(0.85)
CW = W - ML - MR

prs = None
BLANK = None


def new_deck():
    global prs, BLANK
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    BLANK = prs.slide_layouts[6]


def slide(bg=WHITE):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = bg
    return s


def box(s, x, y, w, h, fill=None, line=None, lw=1.0):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    if fill is None: sh.fill.background()
    else: sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None: sh.line.fill.background()
    else: sh.line.color.rgb = line; sh.line.width = Pt(lw)
    sh.shadow.inherit = False
    return sh


def text(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, ls=1.25, sa=0):
    tb = s.shapes.add_textbox(x, y, w, h); tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, it in enumerate(runs):
        body, size, bold, color = it[0], it[1], it[2], it[3]
        fname = it[4] if len(it) > 4 else FONT
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.line_spacing = ls
        if sa: p.space_after = Pt(sa)
        r = p.add_run(); r.text = body
        r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color; r.font.name = fname
    return tb


def header(s, num, title, sub=None):
    text(s, ML, Inches(0.5), CW, Inches(0.3), [(num, 12, False, ACCENT, FONT_M)])
    text(s, ML, Inches(0.8), CW, Inches(0.7), [(title, 27, True, INK)])
    box(s, ML, Inches(1.55), CW, Emu(19050), fill=ACCENT)
    if sub:
        text(s, ML, Inches(1.68), CW, Inches(0.4), [(sub, 13, False, INK_SOFT)])


def footer(s, n, tagline):
    text(s, ML, H - Inches(0.55), CW, Inches(0.24), [(tagline, 9, False, INK_FAINT, FONT_M)])
    text(s, W - MR - Inches(1.0), H - Inches(0.55), Inches(1.0), Inches(0.24),
         [(str(n), 9, False, INK_FAINT, FONT_M)], align=PP_ALIGN.RIGHT)


def table(s, x, y, w, cols, rows, col_w=None, fs=11, row_h=Inches(0.4), head_h=Inches(0.4)):
    n = len(rows) + 1
    t = s.shapes.add_table(n, len(cols), x, y, w, head_h + row_h * len(rows)).table
    if col_w:
        tot = sum(col_w)
        for i, cw in enumerate(col_w): t.columns[i].width = Emu(int(w * cw / tot))
    t.rows[0].height = head_h
    for i in range(1, n): t.rows[i].height = row_h
    for c, label in enumerate(cols):
        cell = t.cell(0, c); cell.fill.solid(); cell.fill.fore_color.rgb = PAPER2
        cell.margin_left = cell.margin_right = Inches(0.09); cell.margin_top = cell.margin_bottom = Inches(0.03)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = cell.text_frame.paragraphs[0]; r = p.add_run(); r.text = label
        r.font.size = Pt(fs); r.font.bold = True; r.font.color.rgb = INK_SOFT; r.font.name = FONT
    for ri, row in enumerate(rows, 1):
        for ci, val in enumerate(row):
            cell = t.cell(ri, ci); cell.fill.solid(); cell.fill.fore_color.rgb = WHITE
            cell.margin_left = cell.margin_right = Inches(0.09); cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]; p.line_spacing = 1.05
            bold = val.startswith("**"); body = val.replace("**", "")
            r = p.add_run(); r.text = body
            r.font.size = Pt(fs); r.font.bold = bold
            r.font.color.rgb = INK if bold else INK_SOFT; r.font.name = FONT
    return t


def stat(s, x, y, w, value, label):
    box(s, x, y, w, Inches(1.2), fill=PAPER2, line=LINE)
    text(s, x + Inches(0.16), y + Inches(0.16), w - Inches(0.32), Inches(0.5),
         [(value, 22, True, ACCENT, FONT_M)])
    text(s, x + Inches(0.16), y + Inches(0.72), w - Inches(0.32), Inches(0.4),
         [(label, 10.5, False, INK_SOFT)], ls=1.1)


def callout(s, x, y, w, h, label, body, tone="accent"):
    fill, edge, lab = ACC_SOFT, ACCENT, ACCENT
    if tone == "warn": fill, edge, lab = GOLD_SOFT, GOLD, GOLD
    elif tone == "good": fill, edge, lab = JADE_SOFT, JADE, JADE
    box(s, x, y, w, h, fill=fill)
    box(s, x, y, Inches(0.05), h, fill=edge)
    text(s, x + Inches(0.22), y + Inches(0.13), w - Inches(0.44), Inches(0.24), [(label, 9, True, lab, FONT_M)])
    text(s, x + Inches(0.22), y + Inches(0.4), w - Inches(0.44), h - Inches(0.52), [(body, 12.5, False, INK)], ls=1.3)


def face(s, x, y, w, h, title, lines, tone):
    fill, tc = (ACC_SOFT, ACCENT) if tone == "out" else (JADE_SOFT, JADE)
    box(s, x, y, w, h, fill=fill, line=LINE)
    text(s, x + Inches(0.2), y + Inches(0.16), w - Inches(0.4), Inches(0.5), [(title, 15, True, tc)])
    text(s, x + Inches(0.2), y + Inches(0.72), w - Inches(0.4), h - Inches(0.9),
         [(l, 12, False, INK_SOFT) for l in lines], ls=1.35, sa=3)


def step(s, x, y, w, h, n, t, d):
    box(s, x, y, w, h, fill=PAPER2, line=LINE)
    text(s, x + Inches(0.14), y + Inches(0.12), w - Inches(0.28), Inches(0.24), [(n, 10, False, ACCENT, FONT_M)])
    text(s, x + Inches(0.14), y + Inches(0.38), w - Inches(0.28), Inches(0.4), [(t, 13, True, INK)], ls=1.05)
    text(s, x + Inches(0.14), y + Inches(0.86), w - Inches(0.28), h - Inches(1.0), [(d, 10.5, False, INK_SOFT)], ls=1.1)


def pic_fit(s, path, x, y, w, h):
    """이미지를 (w,h) 안에 비율 유지로 맞춰 중앙 배치 + 얇은 테두리."""
    pic = s.shapes.add_picture(path, x, y)
    scale = min(w / pic.width, h / pic.height)
    pic.width = int(pic.width * scale); pic.height = int(pic.height * scale)
    pic.left = int(x + (w - pic.width) / 2); pic.top = int(y + (h - pic.height) / 2)
    pic.line.color.rgb = LINE; pic.line.width = Pt(0.75)
    return pic


def screen_grid(s, items, cols, ax, ay, aw, ah):
    """items: [(파일경로, 캡션)]. 격자로 배치, 각 셀 상단에 캡션."""
    import math
    gap = Inches(0.22)
    rows = math.ceil(len(items) / cols)
    cw = (aw - gap * (cols - 1)) / cols
    ch = (ah - gap * (rows - 1)) / rows
    cap_h = Inches(0.3)
    for i, (path, cap) in enumerate(items):
        r, c = divmod(i, cols)
        x = ax + c * (cw + gap); y = ay + r * (ch + gap)
        text(s, x, y, cw, cap_h, [(cap, 10.5, True, ACCENT)], align=PP_ALIGN.CENTER, ls=1.0)
        pic_fit(s, path, x, y + cap_h, cw, ch - cap_h)


def cover(ey, h1, p, tag):
    s = slide(COVER_BG)
    text(s, ML, Inches(2.3), CW, Inches(0.4), [(ey, 13, False, RGBColor(0xCC,0x99,0x88), FONT_M)])
    text(s, ML, Inches(2.75), CW, Inches(1.8), [(h1, 40, True, WHITE)], ls=1.1)
    text(s, ML, Inches(4.5), Inches(8.5), Inches(1.2), [(p, 18, False, RGBColor(0xCC,0xBB,0xBB))], ls=1.3)
    box(s, ML, Inches(6.15), Inches(3.2), Emu(25400), fill=ACCENT)
    text(s, ML, Inches(6.3), CW, Inches(0.4), [(tag, 12, False, RGBColor(0xAA,0x99,0x99), FONT_M)])
    return s


def closing(ey, h1, bullets, tag):
    s = slide(COVER_BG)
    text(s, ML, Inches(0.9), CW, Inches(0.4), [(ey, 13, False, RGBColor(0xCC,0x99,0x88), FONT_M)])
    text(s, ML, Inches(1.35), CW, Inches(1.0), [(h1, 32, True, WHITE)], ls=1.1)
    text(s, ML, Inches(2.7), Inches(11.0), Inches(3.0),
         [("•  " + b, 18, False, RGBColor(0xEE,0xE6,0xE0)) for b in bullets], ls=1.5, sa=10)
    box(s, ML, Inches(6.0), Inches(3.2), Emu(25400), fill=ACCENT)
    text(s, ML, Inches(6.15), CW, Inches(0.8), [(tag, 12, False, RGBColor(0xAA,0x99,0x99), FONT_M)], ls=1.3)
    return s


# ── 콘텐츠 (KO / VI) ────────────────────────────────────────────────
def build(L, outpath):
    new_deck()
    tagline = L["tagline"]
    n = 0

    def foot(s):
        nonlocal n; n += 1; footer(s, n, tagline)

    # 표지
    foot(cover(*L["cover"]))

    # 01 요약
    s = slide(); header(s, "01", *L["s01"]["h"])
    text(s, ML, Inches(2.15), CW, Inches(1.1), [(L["s01"]["lead"], 14, False, INK)], ls=1.35)
    sw = (CW - Inches(0.36)) / 4
    for i, (v, lab) in enumerate(L["s01"]["stats"]):
        stat(s, ML + i * (sw + Inches(0.12)), Inches(3.5), sw, v, lab)
    callout(s, ML, Inches(4.95), CW, Inches(1.7), L["s01"]["co"][0], L["s01"]["co"][1], "good")
    foot(s)

    # 02 배경 (표)
    s = slide(); header(s, "02", *L["s02"]["h"])
    table(s, ML, Inches(2.15), CW, L["s02"]["cols"], L["s02"]["rows"],
          col_w=[26, 24, 26, 24], fs=10.5, row_h=Inches(0.44))
    callout(s, ML, Inches(6.35), CW, Inches(0.75), L["s02"]["co"][0], L["s02"]["co"][1], "good")
    foot(s)

    # 03 구조 (두 얼굴)
    s = slide(); header(s, "03", *L["s03"]["h"])
    fw = (CW - Inches(0.3)) / 2
    face(s, ML, Inches(2.15), fw, Inches(2.2), L["s03"]["out"][0], L["s03"]["out"][1], "out")
    face(s, ML + fw + Inches(0.3), Inches(2.15), fw, Inches(2.2), L["s03"]["in"][0], L["s03"]["in"][1], "in")
    box(s, ML, Inches(4.6), CW, Inches(1.0), fill=INK)
    text(s, ML + Inches(0.24), Inches(4.76), CW - Inches(0.48), Inches(0.7),
         [(L["s03"]["databar"], 12.5, False, WHITE)], ls=1.3)
    callout(s, ML, Inches(5.8), CW, Inches(0.85), L["s03"]["co"][0], L["s03"]["co"][1])
    foot(s)

    # 04 외부, 05 내부, 07 데이터, 08 종목 — 표 슬라이드
    for key, num in (("s04", "04"), ("s05", "05"), ("s07", "07"), ("s08", "08")):
        s = slide(); header(s, num, *L[key]["h"])
        rows = L[key]["rows"]
        table(s, ML, Inches(2.15), CW, L[key]["cols"], rows, col_w=L[key].get("cw", [24, 76]),
              fs=11, row_h=Inches(0.6) if len(rows) <= 6 else Inches(0.52))
        if L[key].get("co"):
            table_h = (Inches(0.4) + (Inches(0.6) if len(rows) <= 6 else Inches(0.52)) * len(rows))
            cy = Inches(2.15) + table_h + Inches(0.12)
            callout(s, ML, cy, CW, Inches(0.8), L[key]["co"][1], L[key]["co"][2], L[key]["co"][0])
        if L[key].get("note"):
            table_h = (Inches(0.4) + (Inches(0.6) if len(rows) <= 6 else Inches(0.52)) * len(rows))
            text(s, ML, Inches(2.15) + table_h + Inches(0.12), CW, Inches(0.5),
                 [(L[key]["note"], 10, False, INK_FAINT)], ls=1.2)
        foot(s)

    # 06 AI 플로우
    s = slide(); header(s, "06", *L["s06"]["h"])
    steps = L["s06"]["steps"]
    stw = (CW - Inches(0.9)) / 4
    for i, (sn, st, sd) in enumerate(steps):
        step(s, ML + i * (stw + Inches(0.3)), Inches(2.15), stw, Inches(1.35), sn, st, sd)
        if i < 3:
            text(s, ML + (i + 1) * stw + i * Inches(0.3) + Inches(0.02), Inches(2.6),
                 Inches(0.28), Inches(0.4), [("→", 18, True, INK_FAINT, FONT_M)], align=PP_ALIGN.CENTER)
    table(s, ML, Inches(3.75), CW, L["s06"]["cols"], L["s06"]["rows"], col_w=[30, 70], fs=11, row_h=Inches(0.5))
    callout(s, ML, Inches(6.15), CW, Inches(0.85), L["s06"]["co"][0], L["s06"]["co"][1], "good")
    foot(s)

    # 09 진행, 10 운영, 11 수익, 12 법 — 표 + 콜아웃
    for key, num in (("s09", "09"), ("s10", "10"), ("s11", "11"), ("s12", "12")):
        s = slide(); header(s, num, *L[key]["h"])
        rows = L[key]["rows"]
        table(s, ML, Inches(2.15), CW, L[key]["cols"], rows, col_w=L[key].get("cw"),
              fs=10.5, row_h=Inches(0.5) if len(rows) > 4 else Inches(0.6))
        if L[key].get("co"):
            rh = (Inches(0.5) if len(rows) > 4 else Inches(0.6))
            cy = Inches(2.15) + Inches(0.4) + rh * len(rows) + Inches(0.12)
            callout(s, ML, cy, CW, Inches(1.15), L[key]["co"][1], L[key]["co"][2], L[key]["co"][0])
        foot(s)

    # 13 리스크
    s = slide(); header(s, "13", *L["s13"]["h"])
    table(s, ML, Inches(2.15), CW, L["s13"]["cols"], L["s13"]["rows"], col_w=[38, 62], fs=11, row_h=Inches(0.62))
    foot(s)

    # 실제 화면 (스크린샷 부록)
    for spec in L.get("screens", []):
        s = slide(); header(s, spec["num"], spec["title"], spec["sub"])
        screen_grid(s, spec["items"], spec["cols"], ML, Inches(2.05), CW, Inches(4.95))
        foot(s)

    # 닫기
    foot(closing(*L["closing"]))

    import os
    os.makedirs("docs", exist_ok=True)
    prs.save(outpath)
    print(f"  saved {outpath} ({n} slides)")


# 콘텐츠 정의는 별도 모듈에서 (가독성)
from _deck_content import KO, VI  # noqa: E402

print("발표자료 생성:")
build(KO, "docs/VSP_사업계획서_KO.pptx")
build(VI, "docs/VSP_사업계획서_VI.pptx")
print("완료.")
