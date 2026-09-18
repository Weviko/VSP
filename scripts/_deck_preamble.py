# -*- coding: utf-8 -*-
"""
VSP 사업계획 발표자료(.pptx) 생성.

레이아웃을 직접 그린다. 파워포인트 기본 템플릿은 자리표시자 위치가 고정되어
한글·베트남어 혼용 제목에서 줄바꿈이 깨지기 때문이다.

  python scripts/make-deck.py
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# ── 팔레트 (사업계획서 웹 문서와 동일) ──────────────────────────────
INK        = RGBColor(0x16, 0x21, 0x2E)
INK_SOFT   = RGBColor(0x46, 0x58, 0x6B)
INK_FAINT  = RGBColor(0x7B, 0x8C, 0xA0)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
PAPER      = RGBColor(0xF6, 0xF8, 0xFA)
PAPER2     = RGBColor(0xED, 0xF1, 0xF5)
LINE       = RGBColor(0xDC, 0xE3, 0xEA)
ACCENT     = RGBColor(0x0E, 0x5C, 0x7A)
ACCENT_SOFT= RGBColor(0xE3, 0xF0, 0xF5)
WARN       = RGBColor(0xC2, 0x41, 0x0C)
WARN_SOFT  = RGBColor(0xFB, 0xED, 0xE6)
GOOD       = RGBColor(0x15, 0x80, 0x3D)
GOOD_SOFT  = RGBColor(0xE6, 0xF2, 0xEA)

# 한글과 베트남어 성조를 모두 지원하는 Windows 기본 글꼴
FONT   = "맑은 고딕"
FONT_M = "Consolas"

W, H = Inches(13.333), Inches(7.5)          # 16:9
ML, MR = Inches(0.85), Inches(0.85)         # 좌우 여백
CONTENT_W = W - ML - MR

prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]


# ── 기본 도구 ────────────────────────────────────────────────────
def slide(bg=WHITE):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = bg
    return s


def box(s, x, y, w, h, fill=None, line=None, line_w=1.0):
    from pptx.enum.shapes import MSO_SHAPE
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(line_w)
    sh.shadow.inherit = False
    return sh


def text(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         line_spacing=1.25, space_after=0):
    """runs: [(문자열, 크기pt, 굵게, 색, 폰트|None), ...] — 각 항목이 한 문단"""
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, item in enumerate(runs):
        body, size, bold, color = item[0], item[1], item[2], item[3]
        fname = item[4] if len(item) > 4 else FONT
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        if space_after:
            p.space_after = Pt(space_after)
        r = p.add_run()
        r.text = body
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        r.font.name = fname
    return tb


def header(s, num, title, sub=None):
    """섹션 번호 + 제목 + 부제. 모든 본문 슬라이드의 상단."""
    text(s, ML, Inches(0.52), CONTENT_W, Inches(0.3),
         [(num, 12, False, ACCENT, FONT_M)])
    text(s, ML, Inches(0.82), CONTENT_W, Inches(0.62),
         [(title, 30, True, INK)])
    box(s, ML, Inches(1.58), CONTENT_W, Emu(12700), fill=ACCENT)
    if sub:
        text(s, ML, Inches(1.74), CONTENT_W, Inches(0.4),
             [(sub, 14, False, INK_SOFT)])


def footer(s, n):
    text(s, ML, H - Inches(0.62), CONTENT_W, Inches(0.24),
         [("Vietnam Sports Platform · 사업계획서 v1.0", 9, False, INK_FAINT, FONT_M)])
    text(s, W - MR - Inches(1.0), H - Inches(0.62), Inches(1.0), Inches(0.24),
         [(str(n), 9, False, INK_FAINT, FONT_M)], align=PP_ALIGN.RIGHT)


def table(s, x, y, w, cols, rows, col_w=None, head_fill=PAPER2,
          fs=11, row_h=Inches(0.36), head_h=Inches(0.38)):
    """가벼운 표. python-pptx 기본 표 스타일은 색이 강해 직접 그린다."""
    n_rows = len(rows) + 1
    shape = s.shapes.add_table(n_rows, len(cols), x, y, w,
                               head_h + row_h * len(rows)).table
    if col_w:
        total = sum(col_w)
        for i, cw in enumerate(col_w):
            shape.columns[i].width = Emu(int(w * cw / total))
    shape.rows[0].height = head_h
    for i in range(1, n_rows):
        shape.rows[i].height = row_h

    for c, label in enumerate(cols):
        cell = shape.cell(0, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = head_fill
        cell.margin_left = cell.margin_right = Inches(0.09)
        cell.margin_top = cell.margin_bottom = Inches(0.04)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = cell.text_frame.paragraphs[0]
        r = p.add_run(); r.text = label
        r.font.size = Pt(fs); r.font.bold = True
        r.font.color.rgb = INK_SOFT; r.font.name = FONT

    for ri, row in enumerate(rows, start=1):
        for ci, val in enumerate(row):
            cell = shape.cell(ri, ci)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE
            cell.margin_left = cell.margin_right = Inches(0.09)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            p.line_spacing = 1.1
            bold = val.startswith("**")
            body = val.replace("**", "")
            r = p.add_run(); r.text = body
            r.font.size = Pt(fs); r.font.bold = bold
            r.font.color.rgb = INK if bold else INK_SOFT
            r.font.name = FONT
    return shape


def stat(s, x, y, w, value, label, color=ACCENT):
    box(s, x, y, w, Inches(1.15), fill=WHITE, line=LINE)
    text(s, x + Inches(0.16), y + Inches(0.16), w - Inches(0.32), Inches(0.48),
         [(value, 26, True, color, FONT_M)])
    text(s, x + Inches(0.16), y + Inches(0.68), w - Inches(0.32), Inches(0.34),
         [(label, 11, False, INK_SOFT)])


def callout(s, x, y, w, h, label, body, tone="accent"):
    fill, edge, lab = ACCENT_SOFT, ACCENT, ACCENT
    if tone == "warn":
        fill, edge, lab = WARN_SOFT, WARN, WARN
    elif tone == "good":
        fill, edge, lab = GOOD_SOFT, GOOD, GOOD
    box(s, x, y, w, h, fill=fill)
    box(s, x, y, Inches(0.045), h, fill=edge)
    text(s, x + Inches(0.24), y + Inches(0.16), w - Inches(0.48), Inches(0.24),
         [(label, 9, True, lab, FONT_M)])
    text(s, x + Inches(0.24), y + Inches(0.44), w - Inches(0.48), h - Inches(0.6),
         [(body, 13, False, INK)], line_spacing=1.35)


def mono_block(s, x, y, w, h, lines):
    box(s, x, y, w, h, fill=PAPER, line=LINE)
    text(s, x + Inches(0.22), y + Inches(0.18), w - Inches(0.44), h - Inches(0.36),
         [(l, 11, False, INK_SOFT, FONT_M) for l in lines], line_spacing=1.3)


